"""Đọc/ghi dữ liệu: thư mục bài làm (.txt/.docx/.pdf), CSV, Excel, JSONL."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

import pandas as pd

TEXT_EXT = {".txt", ".md"}
DOC_EXT = {".docx"}
PDF_EXT = {".pdf"}


def read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as e:
        raise ImportError("Cần cài python-docx: pip install 'viegrader[docs]'") from e
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def read_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as e:
        raise ImportError("Cần cài pdfplumber: pip install 'viegrader[docs]'") from e
    out = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            out.append(page.extract_text() or "")
    return "\n".join(out)


def read_file(path: str | Path) -> str:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in TEXT_EXT:
        return p.read_text(encoding="utf-8", errors="replace")
    if ext in DOC_EXT:
        return read_docx(p)
    if ext in PDF_EXT:
        return read_pdf(p)
    raise ValueError(f"Định dạng không hỗ trợ: {ext}")


def read_folder(folder: str | Path, recursive: bool = True) -> pd.DataFrame:
    """Đọc thư mục bài làm. Tên file được dùng làm essay_id."""
    root = Path(folder)
    pattern = "**/*" if recursive else "*"
    rows: List[Dict[str, Any]] = []
    for p in sorted(root.glob(pattern)):
        if not p.is_file() or p.suffix.lower() not in (TEXT_EXT | DOC_EXT | PDF_EXT):
            continue
        try:
            rows.append({
                "essay_id": p.stem,
                "text": read_file(p),
                "source_path": str(p),
                "prompt_id": p.parent.name if p.parent != root else "",
            })
        except Exception as e:
            rows.append({"essay_id": p.stem, "text": "", "source_path": str(p),
                         "read_error": str(e), "prompt_id": ""})
    return pd.DataFrame(rows)


def read_table(path: str | Path) -> pd.DataFrame:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".csv", ".tsv"):
        return pd.read_csv(p, sep="\t" if ext == ".tsv" else ",")
    if ext in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(p)
    if ext == ".jsonl":
        return pd.read_json(p, lines=True)
    if ext == ".json":
        return pd.read_json(p)
    if ext in (".parquet",):
        return pd.read_parquet(p)
    raise ValueError(f"Định dạng bảng không hỗ trợ: {ext}")


def write_table(df: pd.DataFrame, path: str | Path) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    ext = p.suffix.lower()
    if ext == ".csv":
        df.to_csv(p, index=False, encoding="utf-8-sig")   # utf-8-sig để Excel đọc đúng
    elif ext in (".xlsx", ".xlsm"):
        df.to_excel(p, index=False)
    elif ext == ".jsonl":
        df.to_json(p, orient="records", lines=True, force_ascii=False)
    elif ext == ".json":
        df.to_json(p, orient="records", force_ascii=False, indent=2)
    elif ext == ".parquet":
        df.to_parquet(p, index=False)
    else:
        raise ValueError(f"Định dạng xuất không hỗ trợ: {ext}")


def write_text(text: str, path: str | Path) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")


def write_json(obj: Any, path: str | Path) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def load_input(path: str | Path) -> pd.DataFrame:
    """Nhận cả file bảng lẫn thư mục bài làm."""
    p = Path(path)
    return read_folder(p) if p.is_dir() else read_table(p)
