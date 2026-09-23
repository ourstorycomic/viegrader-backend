from __future__ import annotations

import json
import math
import shutil
import tempfile
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import pandas as pd

from .stages import REPORT_STAGES, detect_stage


@dataclass
class StageResult:
    stage: str
    source_name: str
    source_path: str
    kind: str
    rows: int = 0
    columns: int = 0
    summary: str = ""
    table: Optional[pd.DataFrame] = None
    values: Dict[str, Any] = field(default_factory=dict)

    def manifest(self) -> Dict[str, Any]:
        return {
            "stage": self.stage, "stage_name": REPORT_STAGES[self.stage],
            "source_name": self.source_name, "kind": self.kind,
            "rows": self.rows, "columns": self.columns,
            "summary": self.summary, "values": _json_safe(self.values),
        }


def _json_safe(value):
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if hasattr(value, "item"):
        return _json_safe(value.item())
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _read_source(path: Path) -> StageResult:
    stage, suffix = detect_stage(path), path.suffix.lower()
    result = StageResult(stage, path.name, str(path), suffix.lstrip(".") or "file")
    if suffix in {".csv", ".tsv", ".xlsx", ".xls"}:
        if suffix in {".xlsx", ".xls"}:
            table = pd.read_excel(path)
        else:
            table = pd.read_csv(path, sep="\t" if suffix == ".tsv" else ",")
        result.table, result.rows, result.columns = table, len(table), len(table.columns)
        result.summary = f"Bảng gồm {result.rows} dòng và {result.columns} cột."
        numeric = table.select_dtypes(include="number")
        if len(numeric.columns):
            result.values = {
                str(c): {"mean": float(numeric[c].mean()), "min": float(numeric[c].min()),
                         "max": float(numeric[c].max())}
                for c in numeric.columns[:20] if numeric[c].notna().any()
            }
    elif suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list) and data and isinstance(data[0], dict):
            result.table = pd.DataFrame(data)
            result.rows, result.columns = len(result.table), len(result.table.columns)
            result.summary = f"JSON chứa {result.rows} bản ghi."
        elif isinstance(data, dict):
            result.values = data
            result.summary = f"JSON chứa {len(data)} trường kết quả."
        else:
            result.summary = f"JSON kiểu {type(data).__name__}."
            result.values = {"value": data}
    elif suffix in {".txt", ".md", ".log", ".yaml", ".yml"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        result.summary = text[:12000]
        result.values = {"characters": len(text), "lines": len(text.splitlines())}
    else:
        result.summary = f"Artifact nhị phân, kích thước {path.stat().st_size} byte."
        result.values = {"size_bytes": path.stat().st_size}
    return result


class ReportBuilder:
    def __init__(self, project_name="VieGrader", author="", report_title="Báo cáo kết quả các giai đoạn"):
        self.project_name, self.author, self.report_title = project_name, author, report_title
        self.created_at = datetime.now(timezone.utc).isoformat()
        self.results: List[StageResult] = []

    def ingest(self, paths: Iterable[str | Path]) -> "ReportBuilder":
        for raw in paths:
            path = Path(raw)
            if not path.is_file():
                raise FileNotFoundError(path)
            self.results.append(_read_source(path))
        self.results.sort(key=lambda x: list(REPORT_STAGES).index(x.stage))
        return self

    def summary_frame(self) -> pd.DataFrame:
        rows = [{
            "stage": item.stage, "stage_name": REPORT_STAGES[item.stage],
            "source": item.source_name, "type": item.kind,
            "rows": item.rows, "columns": item.columns,
            "status": "Đã có kết quả",
        } for item in self.results]
        present = {r["stage"] for r in rows}
        rows.extend({
            "stage": stage, "stage_name": name, "source": "", "type": "",
            "rows": 0, "columns": 0, "status": "Chưa cung cấp kết quả",
        } for stage, name in REPORT_STAGES.items() if stage != "other" and stage not in present)
        return pd.DataFrame(rows)

    def manifest(self) -> Dict[str, Any]:
        return {
            "schema_version": "1.0", "project_name": self.project_name,
            "report_title": self.report_title, "author": self.author,
            "created_at": self.created_at,
            "stages": [r.manifest() for r in self.results],
        }

    def export_json(self, path: str | Path) -> Path:
        path = Path(path)
        path.write_text(json.dumps(self.manifest(), ensure_ascii=False, indent=2), encoding="utf-8")
        return path

    def export_html(self, path: str | Path) -> Path:
        parts = [
            "<!doctype html><html lang='vi'><head><meta charset='utf-8'>",
            "<style>body{font-family:Arial;max-width:1100px;margin:32px auto;line-height:1.5}"
            "table{border-collapse:collapse;width:100%;font-size:13px}th,td{border:1px solid #bbb;padding:6px}"
            "th{background:#dbeafe}h1,h2{color:#1e3a8a}.meta{color:#555}</style></head><body>",
            f"<h1>{self.report_title}</h1><p class='meta'>Dự án: {self.project_name}<br>"
            f"Tác giả: {self.author or '—'}<br>Ngày tạo: {self.created_at}</p>",
            "<h2>Tổng hợp</h2>", self.summary_frame().to_html(index=False, escape=True),
        ]
        for stage in REPORT_STAGES:
            items = [r for r in self.results if r.stage == stage]
            if not items:
                continue
            parts.append(f"<h2>{REPORT_STAGES[stage]}</h2>")
            for item in items:
                parts += [f"<h3>{item.source_name}</h3>", f"<p>{item.summary}</p>"]
                if item.table is not None:
                    parts.append(item.table.head(100).to_html(index=False, escape=True))
                elif item.values:
                    parts.append(pd.DataFrame([_json_safe(item.values)]).to_html(index=False, escape=True))
        parts.append("</body></html>")
        path = Path(path); path.write_text("\n".join(parts), encoding="utf-8")
        return path

    def export_xlsx(self, path: str | Path) -> Path:
        path = Path(path)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            self.summary_frame().to_excel(writer, sheet_name="Tong_hop", index=False)
            used = {"Tong_hop"}
            for number, item in enumerate(self.results, 1):
                base = f"{number:02d}_{item.stage}"[:31]
                sheet = base
                counter = 2
                while sheet in used:
                    sheet = f"{base[:27]}_{counter}"; counter += 1
                used.add(sheet)
                if item.table is not None:
                    item.table.to_excel(writer, sheet_name=sheet, index=False)
                else:
                    pd.DataFrame([
                        {"source": item.source_name, "summary": item.summary,
                         "values": json.dumps(_json_safe(item.values), ensure_ascii=False)}
                    ]).to_excel(writer, sheet_name=sheet, index=False)
            workbook = writer.book
            for sheet in workbook.worksheets:
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    font = copy(cell.font); font.bold = True; font.color = "FFFFFF"; cell.font = font
                    cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="1E3A8A")
                for col in sheet.columns:
                    letter = col[0].column_letter
                    width = min(50, max(10, max(len(str(c.value or "")) for c in col) + 2))
                    sheet.column_dimensions[letter].width = width
        return path

    def export_docx(self, path: str | Path) -> Path:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches, Pt
        except ImportError as exc:
            raise ImportError("Cài viegrader[reports] để xuất Word.") from exc
        doc = Document()
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.7)
        title = doc.add_heading(self.report_title, 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Dự án: {self.project_name}\nTác giả: {self.author or '—'}\nNgày tạo: {self.created_at}")
        doc.add_heading("Tổng hợp tiến độ", level=1)
        _add_docx_table(doc, self.summary_frame(), 30)
        for stage in REPORT_STAGES:
            items = [r for r in self.results if r.stage == stage]
            if not items:
                continue
            doc.add_heading(REPORT_STAGES[stage], level=1)
            for item in items:
                doc.add_heading(item.source_name, level=2)
                doc.add_paragraph(item.summary or "Không có mô tả.")
                if item.table is not None:
                    _add_docx_table(doc, item.table, 30)
                elif item.values:
                    _add_docx_table(doc, pd.DataFrame([_json_safe(item.values)]), 10)
        styles = doc.styles
        styles["Normal"].font.name = "Times New Roman"; styles["Normal"].font.size = Pt(12)
        path = Path(path); doc.save(path); return path

    def export_pdf(self, path: str | Path) -> Path:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import cm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise ImportError("Cài viegrader[reports] để xuất PDF.") from exc
        font_name = "Helvetica"
        for candidate in (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        ):
            if Path(candidate).exists():
                pdfmetrics.registerFont(TTFont("VieGraderUnicode", candidate))
                font_name = "VieGraderUnicode"
                break
        path = Path(path); styles = getSampleStyleSheet(); story = []
        for style in styles.byName.values():
            style.fontName = font_name
        story += [Paragraph(self.report_title, styles["Title"]), Spacer(1, 0.3 * cm),
                  Paragraph(f"Dự án: {self.project_name}<br/>Tác giả: {self.author or '—'}<br/>Ngày tạo: {self.created_at}", styles["Normal"]),
                  Spacer(1, 0.4 * cm), Paragraph("Tổng hợp tiến độ", styles["Heading1"]),
                  _pdf_table(self.summary_frame(), Table, TableStyle, colors, font_name)]
        for stage in REPORT_STAGES:
            items = [r for r in self.results if r.stage == stage]
            if not items: continue
            story += [PageBreak(), Paragraph(REPORT_STAGES[stage], styles["Heading1"])]
            for item in items:
                story += [Paragraph(item.source_name, styles["Heading2"]),
                          Paragraph((item.summary or "Không có mô tả.")[:4000].replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 0.2 * cm)]
                if item.table is not None:
                    story.append(_pdf_table(item.table.head(25), Table, TableStyle, colors, font_name))
        SimpleDocTemplate(str(path), pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm,
                          topMargin=1*cm, bottomMargin=1*cm).build(story)
        return path


def _add_docx_table(doc, frame: pd.DataFrame, max_rows: int) -> None:
    frame = frame.head(max_rows).fillna("")
    cols = list(frame.columns[:12])
    table = doc.add_table(rows=1, cols=len(cols)); table.style = "Table Grid"
    for i, col in enumerate(cols): table.rows[0].cells[i].text = str(col)
    for _, row in frame[cols].iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(cols): cells[i].text = str(row[col])[:500]


def _pdf_table(frame, Table, TableStyle, colors, font_name="Helvetica"):
    data_frame = frame.head(25).fillna("")
    cols = list(data_frame.columns[:10])
    data = [[str(c)[:35] for c in cols]] + [[str(row[c])[:80] for c in cols] for _, row in data_frame[cols].iterrows()]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return table


def build_report_bundle(
    paths: Sequence[str | Path], project_name="VieGrader", author="",
    formats: Sequence[str] = ("xlsx", "docx", "html", "json"),
    output_dir: Optional[str | Path] = None,
) -> Tuple[Path, pd.DataFrame, List[Path]]:
    if not paths:
        raise ValueError("Cần ít nhất một tệp kết quả.")
    root = Path(output_dir) if output_dir else Path(tempfile.mkdtemp(prefix="viegrader_reports_"))
    root.mkdir(parents=True, exist_ok=True)
    builder = ReportBuilder(project_name=project_name, author=author).ingest(paths)
    outputs = []
    exporters = {
        "json": builder.export_json, "html": builder.export_html,
        "xlsx": builder.export_xlsx, "docx": builder.export_docx,
        "pdf": builder.export_pdf,
    }
    for fmt in formats:
        key = str(fmt).lower().lstrip(".")
        if key not in exporters:
            raise ValueError(f"Định dạng không hỗ trợ: {fmt}")
        outputs.append(exporters[key](root / f"bao_cao_cac_giai_doan.{key}"))
    zip_base = root.parent / f"{root.name}_BaoCao_VieGrader"
    zip_path = Path(shutil.make_archive(str(zip_base), "zip", root_dir=root,
                                        base_dir="."))
    return zip_path, builder.summary_frame(), outputs
